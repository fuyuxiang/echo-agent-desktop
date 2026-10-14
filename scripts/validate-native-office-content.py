#!/usr/bin/env python3
"""Independently open documents produced by native-office-review.mjs.

This script only reads existing receipts and artifacts. It does not create
documents, launch a Worker, replace IPC, or turn fixture results into native
acceptance. Use it after the real desktop Office MCP review has succeeded.
"""

import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import stat
import sys
from xml.etree import ElementTree
from zipfile import ZipFile


MAX_FILE_BYTES = 50 * 1024 * 1024
TITLE = "原生中文办公验收"
PARENT = "父任务：核对成果"
CHILDREN = ["子任务甲：确认中文与表格", "子任务乙：确认代码换行"]
FOLLOWING = "后续任务：保留第一次生成的文件"
CODE_LINES = ["第一行代码", "第二行代码"]
TABLE_ROWS = [
    ["项目", "编号", "状态", "表达式", "金额"],
    ["中文任务甲", "001234", "已核对", "$a<b>c$", "$5"],
    ["中文任务乙", "12345678901234567890", "待复核", r"$\{x\}$", "$10"],
]


def require(condition, message):
    if not condition:
        raise ValueError(message)


def compact(text):
    return re.sub(r"\s+", "", text)


def check_full_text(text):
    normalized = compact(text)
    required = [TITLE, PARENT, *CHILDREN, FOLLOWING, *CODE_LINES]
    required.extend(["a < b > c", "$a<b>c$", r"$\{x\}$", "$$", r"\{x\}", "x<y>z", "flowchart LR", "A[甲] --> B[乙]"])
    required.extend([r"\(a<b>c\)", r"\[\{x\}\]"])
    required.extend(["$5 / $10", "office affinity efficient final_file"])
    required.extend(value for row in TABLE_ROWS for value in row)
    for value in required:
        require(compact(value) in normalized, f"Saved document is missing expected content: {value}")
    offsets = [normalized.index(compact(value)) for value in [PARENT, *CHILDREN, FOLLOWING]]
    require(offsets == sorted(offsets), "Saved document changed the nested-list item order")


def check_docx(path):
    from docx import Document

    document = Document(path)
    paragraphs = document.paragraphs
    table_rows = [[cell.text for cell in row.cells] for table in document.tables for row in table.rows]
    text = "\n".join([paragraph.text for paragraph in paragraphs] + [" | ".join(row) for row in table_rows])
    check_full_text(text)
    require(TABLE_ROWS == table_rows, "Word table headers, Chinese cells, or text identifiers differ")
    list_paragraphs = []
    for value in [PARENT, *CHILDREN, FOLLOWING]:
        found = [paragraph for paragraph in paragraphs if value in paragraph.text]
        require(len(found) == 1, "Word nested-list item is missing or duplicated")
        list_paragraphs.append(found[0])
    indents = [paragraph.paragraph_format.left_indent for paragraph in list_paragraphs]
    require(all(indent is not None for indent in indents), "Word list items have no explicit hierarchy indentation")
    require(indents[1] > indents[0] and indents[2] > indents[0] and indents[3] == indents[0],
            "Word nested-list indentation does not retain parent/child hierarchy")
    code = [paragraph for paragraph in paragraphs if all(line in paragraph.text for line in CODE_LINES)]
    require(len(code) == 1 and code[0].text.splitlines() == CODE_LINES,
            "Word code lines are missing, combined, or reordered")
    with ZipFile(path) as archive:
        xml = ElementTree.fromstring(archive.read("word/document.xml"))
    namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    require(bool(xml.findall(".//w:br", namespaces)), "Word code has no actual OOXML line break")
    return {
        "openedWith": "python-docx + OOXML",
        "paragraphs": len(paragraphs), "tables": len(document.tables),
        "listIndentTwips": [indent.twips for indent in indents],
        "tableRows": table_rows, "codeLineBreaks": True, "chineseContent": True,
    }


def check_xlsx(path):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        require(len(workbook.worksheets) == 1, "Excel output does not contain the expected single Markdown table")
        sheet = workbook.worksheets[0]
        rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
        require(rows == TABLE_ROWS, "Excel table content or identifier precision differs")
        require(sheet.cell(2, 2).data_type == "s" and sheet.cell(3, 2).data_type == "s",
                "Excel leading-zero or long identifier was converted into a number")
        return {
            "openedWith": "openpyxl", "worksheets": workbook.sheetnames,
            "tableRows": rows, "identifiersStoredAsText": True, "chineseContent": True,
            "scope": "Excel exports Markdown tables when the source contains tables",
        }
    finally:
        workbook.close()


def check_pdf(path, render_directory):
    import fitz

    with fitz.open(path) as document:
        require(document.is_pdf and not document.is_encrypted and len(document) > 0,
                "Generated PDF is not an open, unencrypted document")
        check_full_text("\n".join(page.get_text() for page in document))
        found = {value: [] for value in CODE_LINES}
        for page_index, page in enumerate(document):
            for block in page.get_text("dict")["blocks"]:
                for line in block.get("lines", []):
                    text = compact("".join(span["text"] for span in line["spans"]))
                    for value in CODE_LINES:
                        if value in text:
                            found[value].append((page_index, line["bbox"][1], line["bbox"][3]))
        require(all(len(found[value]) == 1 for value in CODE_LINES), "PDF code lines cannot be found as separate text lines")
        first, second = found[CODE_LINES[0]][0], found[CODE_LINES[1]][0]
        require(second[0] > first[0] or (second[0] == first[0] and second[1] > first[1] + 1),
                "PDF code lines overlap or are reordered")
        images = []
        if render_directory:
            render_directory.mkdir(parents=True, exist_ok=True)
            for page_index, page in enumerate(document):
                image = render_directory / f"native-office-pdf-page-{page_index + 1}.png"
                page.get_pixmap(matrix=fitz.Matrix(1.25, 1.25)).save(image)
                images.append(str(image.resolve()))
        return {
            "openedWith": "PyMuPDF", "pages": len(document),
            "chineseContent": True, "codeLinePositions": found, "renderedPages": images,
        }


def check_pptx(path):
    from pptx import Presentation

    presentation = Presentation(path)
    frames = []
    table_text = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                frames.append(shape.text)
            if shape.has_table:
                table_text.extend(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
    check_full_text("\n".join(frames + table_text))
    all_lines = [line for frame in frames for line in frame.splitlines()]
    list_lines = []
    for value in [PARENT, *CHILDREN, FOLLOWING]:
        matching = [line for line in all_lines if value in line]
        require(len(matching) == 1, "PowerPoint nested-list item is missing or duplicated")
        list_lines.append(matching[0])
    indents = [len(re.match(r"^\s*", line).group()) for line in list_lines]
    require(indents[1] > indents[0] and indents[2] > indents[0] and indents[3] == indents[0],
            "PowerPoint text lost the nested-list indentation")
    code_frames = [frame.splitlines() for frame in frames if all(line in frame for line in CODE_LINES)]
    require(len(code_frames) == 1, "PowerPoint code lines are missing or split into unrelated frames")
    lines = code_frames[0]
    require(CODE_LINES[0] in lines and CODE_LINES[1] in lines
            and lines.index(CODE_LINES[0]) < lines.index(CODE_LINES[1]),
            "PowerPoint code lines are combined or reordered")
    return {
        "openedWith": "python-pptx", "slides": len(presentation.slides),
        "textFrames": len(frames), "listLeadingSpaces": indents,
        "chineseContent": True, "codeLineBreaks": True,
        "scope": "Markdown table rows are preserved as presentation text",
    }


def verify_receipt(receipt, directory):
    require(isinstance(receipt, dict) and receipt.get("format") in {"docx", "pdf", "xlsx", "pptx"},
            "Invalid native Office receipt format")
    require(isinstance(receipt.get("path"), str), "Missing native Office artifact path")
    path = Path(receipt["path"])
    require(path.is_absolute() and path.parent == directory and path.resolve(strict=True) == path,
            "Native Office artifact is outside its ordinary authorized output directory")
    metadata = path.lstat()
    require(stat.S_ISREG(metadata.st_mode) and not path.is_symlink(), "Native Office artifact is not an ordinary file")
    require(type(receipt.get("byteSize")) is int and 0 < metadata.st_size <= MAX_FILE_BYTES
            and metadata.st_size == receipt["byteSize"], "Native Office artifact byte size differs")
    require(path.suffix == "." + receipt["format"], "Native Office artifact extension differs")
    data = path.read_bytes()
    digest = hashlib.sha256(data).hexdigest()
    require(digest == receipt.get("sha256"), "Native Office artifact SHA-256 differs")
    require(data.startswith(b"%PDF-") if receipt["format"] == "pdf" else data.startswith(b"PK\x03\x04"),
            "Native Office artifact file signature differs")
    return path


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("receipt_json", type=Path, help="Actual successful native-office.json from the running desktop review")
    parser.add_argument("--home", type=Path, help="Optional explicit isolated home; otherwise inferred from workspace/workspaces/default")
    parser.add_argument("--output", type=Path, help="Result JSON; default next to native-office.json")
    parser.add_argument("--render-dir", type=Path, help="Optionally render all real PDF pages into this evidence directory")
    args = parser.parse_args()
    output = args.output or args.receipt_json.with_name("native-office-content.json")
    result = {
        "environment": "independent parser verification of existing native Office MCP artifacts",
        "passed": False, "sourceReceipt": str(args.receipt_json.resolve()), "documents": [],
    }
    try:
        source = json.loads(args.receipt_json.read_text())
        require(source.get("passed") is True, "Native Office MCP generation did not pass; content validation cannot replace it")
        require(source.get("environment") == "native desktop Office MCP + production Worker + real files",
                "Receipt provenance differs from the native Office review")
        receipts = source.get("receipts")
        require(isinstance(receipts, list) and len(receipts) == 5, "Expected four formats and a repeated-title DOCX")
        require(Counter(item.get("format") for item in receipts) == Counter({"docx": 2, "pdf": 1, "xlsx": 1, "pptx": 1}),
                "Native Office receipt list has incomplete format coverage")
        require(isinstance(source.get("workspace"), str), "Native Office receipt has no workspace")
        workspace = Path(source["workspace"])
        home = (args.home or workspace.parent.parent).resolve(strict=True)
        require(workspace.is_absolute() and workspace.resolve(strict=True) == home / "workspaces" / "default",
                "Native Office receipt workspace differs from the isolated default workspace")
        directory = workspace / "EchoAgent成果"
        require(directory.is_dir() and not directory.is_symlink() and directory.resolve(strict=True) == directory,
                "Native Office output directory is not an ordinary workspace child")
        require(len({item.get("path") for item in receipts}) == 5, "Repeated document title replaced an earlier artifact")
        for receipt in receipts:
            path = verify_receipt(receipt, directory)
            if receipt["format"] == "docx":
                details = check_docx(path)
            elif receipt["format"] == "xlsx":
                details = check_xlsx(path)
            elif receipt["format"] == "pdf":
                details = check_pdf(path, args.render_dir)
            else:
                details = check_pptx(path)
            result["documents"].append({
                "format": receipt["format"], "path": str(path),
                "byteSize": receipt["byteSize"], "sha256": receipt["sha256"], **details,
            })
        result["passed"] = True
    except Exception as error:
        result["error"] = str(error)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    sys.exit(main())
